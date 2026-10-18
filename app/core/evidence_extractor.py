from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.hashing import sha256_bytes
from app.core.models import Evidence, EvidenceStatus, Source, SourceKind
from app.core.source_index import ID_SEPARATOR, SourceIndex, output_is_outside_source

EVIDENCE_ID_LENGTH = 16
EVIDENCE_FILE_NAME = "evidence.json"
BLOCKED_OUTPUT_INSIDE_SOURCE = "blocked:output_inside_source_root"

SHEET_PRODUCT_DATA = "商品参数"
SHEET_REVIEWS = "用户评价"
SHEET_HISTORY = "历史内容"
SHEET_BRAND_MEMORY = "品牌记忆"
EXPECTED_XLSX_SHEETS = (
    SHEET_PRODUCT_DATA,
    SHEET_REVIEWS,
    SHEET_HISTORY,
    SHEET_BRAND_MEMORY,
)

HEADER_PRODUCT_DATA = "字段"
HEADER_REVIEWS = "编号"
HEADER_HISTORY = "日期"
HEADER_BRAND_MEMORY = "项目"

KEY_VALUE_PATTERN = re.compile(r"^([^：:]+)[：:]\s*(.+)$")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+)$")
BULLET_PATTERN = re.compile(r"^[-*]\s+(.+)$")

TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gbk")


def evidence_id_of(
    product_id: str,
    source_id: str,
    field_id: str,
    source_locator: str,
) -> str:
    joined = (
        f"{product_id}{ID_SEPARATOR}{source_id}{ID_SEPARATOR}"
        f"{field_id}{ID_SEPARATOR}{source_locator}"
    )
    return "evd_" + sha256_bytes(joined.encode("utf-8"))[:EVIDENCE_ID_LENGTH]


def cell_locator(file_name: str, sheet_name: str, row: int, column: int) -> str:
    return f"{file_name}!{sheet_name}!{_column_letter(column)}{row}"


def line_locator(file_name: str, line_number: int) -> str:
    return f"{file_name}:L{line_number}"


def heading_locator(file_name: str, heading: str, line_number: int) -> str:
    return f"{file_name}:{heading}:L{line_number}"


@dataclass(frozen=True)
class SkippedExtraction:
    source_id: str
    relative_path: str
    reason: str
    detail: str | None = None


@dataclass(frozen=True)
class EvidenceExtractionReport:
    product_id: str
    library_id: str
    evidences: tuple[Evidence, ...]
    skipped: tuple[SkippedExtraction, ...]

    def by_field_id(self, field_id: str) -> tuple[Evidence, ...]:
        return tuple(item for item in self.evidences if item.field_id == field_id)

    def values_for_field(self, field_id: str) -> tuple[Any, ...]:
        return tuple(item.value for item in self.by_field_id(field_id))

    def to_dict(self) -> dict[str, Any]:
        return {
            "product_id": self.product_id,
            "library_id": self.library_id,
            "evidence_count": len(self.evidences),
            "skipped_count": len(self.skipped),
            "evidences": [evidence_to_dict(item) for item in self.evidences],
            "skipped": [
                {
                    "source_id": item.source_id,
                    "relative_path": item.relative_path,
                    "reason": item.reason,
                    "detail": item.detail,
                }
                for item in self.skipped
            ],
        }

    def write_json(self, output_root: Path, library_root: Path) -> list[str]:
        if not output_is_outside_source(output_root, library_root):
            return [BLOCKED_OUTPUT_INSIDE_SOURCE]
        destination = output_root / self.product_id / EVIDENCE_FILE_NAME
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            temporary = destination.with_name(destination.name + ".tmp")
            temporary.write_text(
                json.dumps(self.to_dict(), ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            os.replace(temporary, destination)
        except OSError as error:
            return [f"write_failed:{destination.as_posix()}:{type(error).__name__}"]
        return []


def evidence_to_dict(evidence: Evidence) -> dict[str, Any]:
    return {
        "evidence_id": evidence.evidence_id,
        "product_id": evidence.product_id,
        "source_id": evidence.source_id,
        "field_id": evidence.field_id,
        "raw_text": evidence.raw_text,
        "value": evidence.value,
        "scope": evidence.scope,
        "unit": evidence.unit,
        "source_locator": evidence.source_locator,
        "status": evidence.status.value,
        "source_hash": evidence.source_hash,
        "invalidated_reason": evidence.invalidated_reason,
    }


def extract_evidence_from_index(
    index: SourceIndex,
    library_root: Path | None = None,
) -> EvidenceExtractionReport:
    root = Path(library_root or index.library_root)
    evidences: list[Evidence] = []
    skipped: list[SkippedExtraction] = []

    for source in index.sources:
        path = root / source.relative_path
        try:
            extracted = extract_evidence_for_source(source, path)
        except OSError as error:
            skipped.append(
                SkippedExtraction(
                    source_id=source.source_id,
                    relative_path=source.relative_path,
                    reason="unreadable_file",
                    detail=type(error).__name__,
                )
            )
            continue
        if extracted is None:
            skipped.append(
                SkippedExtraction(
                    source_id=source.source_id,
                    relative_path=source.relative_path,
                    reason="unsupported_source_kind",
                    detail=source.source_kind.value,
                )
            )
            continue
        evidences.extend(extracted)

    return EvidenceExtractionReport(
        product_id=index.product_id,
        library_id=index.library_id,
        evidences=tuple(evidences),
        skipped=tuple(skipped),
    )


def extract_evidence_for_source(source: Source, path: Path) -> list[Evidence] | None:
    if not path.is_file():
        raise OSError("missing_file")

    if source.file_type == "xlsx":
        return _extract_xlsx(source, path)
    if source.file_type in {"txt", "md"}:
        text = _read_text(path)
        if source.file_type == "txt":
            return _extract_plain_text(source, text)
        return _extract_markdown(source, text)
    if source.source_kind in {SourceKind.IMAGE, SourceKind.VIDEO}:
        return None
    if source.file_type in {"csv"}:
        return _extract_plain_text(source, _read_text(path))
    return None


def _extract_xlsx(source: Source, path: Path) -> list[Evidence]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    evidences: list[Evidence] = []
    missing_sheets = [name for name in EXPECTED_XLSX_SHEETS if name not in workbook.sheetnames]
    if missing_sheets:
        workbook.close()
        raise ValueError(f"missing_sheets:{','.join(missing_sheets)}")

    evidences.extend(_extract_product_data_sheet(source, workbook[SHEET_PRODUCT_DATA]))
    evidences.extend(_extract_table_sheet(
        source,
        workbook[SHEET_REVIEWS],
        sheet_name=SHEET_REVIEWS,
        field_prefix="review",
        row_key_column=1,
    ))
    evidences.extend(_extract_table_sheet(
        source,
        workbook[SHEET_HISTORY],
        sheet_name=SHEET_HISTORY,
        field_prefix="history",
        row_key_column=1,
    ))
    evidences.extend(_extract_table_sheet(
        source,
        workbook[SHEET_BRAND_MEMORY],
        sheet_name=SHEET_BRAND_MEMORY,
        field_prefix="brand",
        row_key_column=1,
    ))
    workbook.close()
    return evidences


def _extract_product_data_sheet(source: Source, worksheet: Worksheet) -> list[Evidence]:
    header_row, headers = _find_header_row(worksheet, HEADER_PRODUCT_DATA)
    evidences: list[Evidence] = []
    for row_number, row in _iter_data_rows(worksheet, header_row + 1):
        if not row:
            continue
        field_name = _cell_text(row[0])
        if not field_name:
            continue
        value = row[1] if len(row) > 1 else None
        scope = _cell_text(row[2]) if len(row) > 2 else None
        hint = _cell_text(row[3]) if len(row) > 3 else ""
        raw_text = _format_row_text(headers, row)
        field_id = f"excel.{SHEET_PRODUCT_DATA}.{field_name}"
        locator = cell_locator(source.file_name, SHEET_PRODUCT_DATA, row_number, 2)
        evidences.append(
            _make_evidence(
                source=source,
                field_id=field_id,
                raw_text=raw_text,
                value=value,
                scope=scope,
                source_locator=locator,
            )
        )
        if hint:
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id=f"excel.{SHEET_PRODUCT_DATA}.{field_name}.数据使用提示",
                    raw_text=f"{field_name}：{hint}",
                    value=hint,
                    scope=scope,
                    source_locator=cell_locator(
                        source.file_name,
                        SHEET_PRODUCT_DATA,
                        row_number,
                        4,
                    ),
                )
            )
    return evidences


def _extract_table_sheet(
    source: Source,
    worksheet: Worksheet,
    *,
    sheet_name: str,
    field_prefix: str,
    row_key_column: int,
) -> list[Evidence]:
    header_marker = {
        SHEET_REVIEWS: HEADER_REVIEWS,
        SHEET_HISTORY: HEADER_HISTORY,
        SHEET_BRAND_MEMORY: HEADER_BRAND_MEMORY,
    }[sheet_name]
    header_row, headers = _find_header_row(worksheet, header_marker)
    evidences: list[Evidence] = []
    for row_number, row in _iter_data_rows(worksheet, header_row + 1):
        if not row or len(row) < row_key_column:
            continue
        row_key = _cell_text(row[row_key_column - 1])
        if not row_key:
            continue
        for column_index, header in enumerate(headers, start=1):
            if column_index > len(row):
                continue
            value = row[column_index - 1]
            if _is_empty(value):
                continue
            header_name = header or f"column_{column_index}"
            field_id = f"excel.{sheet_name}.{row_key}.{header_name}"
            raw_text = f"{header_name}：{_cell_text(value)}"
            locator = cell_locator(
                source.file_name,
                sheet_name,
                row_number,
                column_index,
            )
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id=field_id,
                    raw_text=raw_text,
                    value=value,
                    scope=None,
                    source_locator=locator,
                )
            )
    return evidences


def _extract_plain_text(source: Source, text: str) -> list[Evidence]:
    evidences: list[Evidence] = []
    current_heading = source.file_name
    for line_number, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        heading_match = HEADING_PATTERN.match(stripped)
        if heading_match:
            current_heading = heading_match.group(2).strip()
            continue
        key_value_match = KEY_VALUE_PATTERN.match(stripped)
        if key_value_match:
            field_name = key_value_match.group(1).strip()
            value_text = key_value_match.group(2).strip()
            field_id = f"text.{_slug_field(field_name)}"
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id=field_id,
                    raw_text=stripped,
                    value=value_text,
                    scope=current_heading,
                    source_locator=line_locator(source.file_name, line_number),
                )
            )
            continue
        if re.match(r"^\d{4}-\d{2}-\d{2}\b", stripped):
            field_id = f"text.history_entry.{stripped.split()[0]}"
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id=field_id,
                    raw_text=stripped,
                    value=stripped,
                    scope=current_heading,
                    source_locator=line_locator(source.file_name, line_number),
                )
            )
            continue
        if stripped.startswith("结论："):
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id="text.history_summary",
                    raw_text=stripped,
                    value=stripped.removeprefix("结论：").strip(),
                    scope=current_heading,
                    source_locator=line_locator(source.file_name, line_number),
                )
            )
    return evidences


def _extract_markdown(source: Source, text: str) -> list[Evidence]:
    evidences: list[Evidence] = []
    heading_stack: list[str] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        heading_match = HEADING_PATTERN.match(stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            heading_stack = heading_stack[: level - 1]
            heading_stack.append(title)
            continue
        current_heading = " / ".join(heading_stack) if heading_stack else source.file_name
        bullet_match = BULLET_PATTERN.match(stripped)
        content = bullet_match.group(1).strip() if bullet_match else stripped
        key_value_match = KEY_VALUE_PATTERN.match(content)
        if key_value_match:
            field_name = key_value_match.group(1).strip()
            value_text = key_value_match.group(2).strip()
            field_id = f"md.{_slug_field(current_heading)}.{_slug_field(field_name)}"
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id=field_id,
                    raw_text=stripped,
                    value=value_text,
                    scope=current_heading,
                    source_locator=heading_locator(
                        source.file_name,
                        current_heading,
                        line_number,
                    ),
                )
            )
            continue
        if bullet_match:
            field_id = f"md.{_slug_field(current_heading)}.bullet_{line_number}"
            evidences.append(
                _make_evidence(
                    source=source,
                    field_id=field_id,
                    raw_text=stripped,
                    value=content,
                    scope=current_heading,
                    source_locator=heading_locator(
                        source.file_name,
                        current_heading,
                        line_number,
                    ),
                )
            )
    return evidences


def _make_evidence(
    *,
    source: Source,
    field_id: str,
    raw_text: str,
    value: Any,
    scope: str | None,
    source_locator: str,
    unit: str | None = None,
) -> Evidence:
    return Evidence(
        evidence_id=evidence_id_of(
            source.product_id,
            source.source_id,
            field_id,
            source_locator,
        ),
        product_id=source.product_id,
        source_id=source.source_id,
        field_id=field_id,
        raw_text=raw_text,
        value=value,
        scope=scope,
        unit=unit,
        source_locator=source_locator,
        status=EvidenceStatus.PENDING,
        source_hash=source.file_hash,
    )


def _find_header_row(worksheet: Worksheet, marker: str) -> tuple[int, tuple[str, ...]]:
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if not row:
            continue
        if _cell_text(row[0]) == marker:
            headers = tuple(_cell_text(cell) for cell in row)
            return row_number, headers
    raise ValueError(f"header_not_found:{marker}")


def _iter_data_rows(
    worksheet: Worksheet,
    start_row: int,
) -> Iterable[tuple[int, tuple[Any, ...]]]:
    for offset, row in enumerate(worksheet.iter_rows(min_row=start_row, values_only=True)):
        if not any(not _is_empty(cell) for cell in row):
            continue
        yield start_row + offset, row


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _column_letter(column: int) -> str:
    letters = ""
    current = column
    while current:
        current, remainder = divmod(current - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _format_row_text(headers: tuple[str, ...], row: tuple[Any, ...]) -> str:
    parts: list[str] = []
    for index, header in enumerate(headers):
        if not header:
            continue
        if index >= len(row) or _is_empty(row[index]):
            continue
        parts.append(f"{header}：{_cell_text(row[index])}")
    return " | ".join(parts)


def _slug_field(value: str) -> str:
    cleaned = re.sub(r"\s+", "_", value.strip())
    cleaned = cleaned.replace("/", "_")
    return cleaned or "field"
