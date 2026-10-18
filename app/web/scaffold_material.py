from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl


def _has_workbook(material_root: Path) -> bool:
    return any(material_root.rglob("*.xlsx")) or any(material_root.rglob("*.xls"))


def ensure_material_brief_xlsx(material_root: Path, product_form: dict[str, Any] | None = None) -> Path | None:
    """Create a minimal Excel brief when the user only uploaded images."""
    if _has_workbook(material_root):
        return None

    form = product_form or {}
    destination = material_root / "商品资料.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "商品参数"
    sheet.append(["商品参数"])
    sheet.append(["由上传图片自动生成"])
    sheet.append([])
    sheet.append(["字段", "值", "分类", "数据使用提示"])

    rows = [
        ("商品名称", form.get("product_name") or "未命名商品", "商品", ""),
        ("商品类别", form.get("product_category") or "商品", "商品", ""),
        ("建议价格", form.get("product_price") or "", "价格", ""),
        ("主要场景", form.get("product_desc") or "日常使用", "场景", ""),
        ("品牌名称", form.get("product_brand") or "", "品牌", ""),
        ("材质", "待补充", "商品参数", "可结合图片理解"),
    ]
    for row in rows:
        sheet.append(list(row))

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["由上传图片自动生成"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["U01", 5, "图片素材已上传，待 AI 结合图像理解商品。", "正向", "图片", "参考图"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["由上传图片自动生成"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-01-01", form.get("product_name") or "当前商品", "参考图创作", "待生成", "低"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["由上传图片自动生成"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["品牌调性", form.get("product_brand") or "清新自然", "结合参考图保持统一"])

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination
