from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.core.evidence_extractor import (
    EVIDENCE_FILE_NAME,
    cell_locator,
    evidence_id_of,
    extract_evidence_for_source,
    extract_evidence_from_index,
    line_locator,
)
from app.core.models import EvidenceStatus, Source, SourceKind, Status
from app.core.source_index import source_id_of
from app.core.source_scanner import scan_library


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    product_sheet = workbook.active
    product_sheet.title = "商品参数"
    product_sheet.append(["界线 直筒牛仔裤 01 中靛蓝 商品参数"])
    product_sheet.append(["演示数据"])
    product_sheet.append([])
    product_sheet.append(["字段", "值", "分类", "数据使用提示"])
    product_sheet.append(["建议价格", 329, "价格", "展示为 329 元"])
    product_sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    product_sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    product_sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "倾向", "标签", "使用场景"])
    reviews.append(["J01", 5, "腰头有支撑，坐下不勒", "好评", "腰部", "办公室"])
    reviews.append(["J02", 5, "直筒不贴腿", "好评", "直筒", "通勤"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达提醒"])
    history.append(["2026-04-18", "晴日直筒裤", "小个子直筒牛仔裤", "显瘦 小个子", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤 闭眼入", "禁入文案"])
    workbook.save(path)


def _source(tmp_path: Path, relative_path: str, **overrides) -> Source:
    payload = {
        "source_id": source_id_of("lib_test000000", relative_path),
        "product_id": "prod_test000001",
        "file_name": Path(relative_path).name,
        "relative_path": relative_path,
        "source_kind": SourceKind.PRODUCT_DATA,
        "file_type": Path(relative_path).suffix.lstrip(".").lower(),
        "file_hash": "hash",
        "parse_status": Status.READY,
        "associated_product_id": "prod_test000001",
        "source_folder": None,
    }
    payload.update(overrides)
    return Source(**payload)


def test_evidence_id_is_stable_and_does_not_use_absolute_paths() -> None:
    first = evidence_id_of("prod_a", "src_a", "excel.商品参数.建议价格", "Demo.xlsx!商品参数!B5")
    second = evidence_id_of("prod_a", "src_a", "excel.商品参数.建议价格", "Demo.xlsx!商品参数!B5")
    third = evidence_id_of("prod_b", "src_a", "excel.商品参数.建议价格", "Demo.xlsx!商品参数!B5")

    assert first == second
    assert first != third
    assert "E:" not in first
    assert "demo" not in first.lower()


def test_xlsx_extractor_emits_locators_and_values(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "Demo经营资料.xlsx"
    _write_xlsx(xlsx_path)
    source = _source(tmp_path, "Demo经营资料.xlsx")

    evidences = extract_evidence_for_source(source, xlsx_path)

    price = [item for item in evidences if item.field_id == "excel.商品参数.建议价格"]
    assert len(price) == 1
    assert price[0].value == 329
    assert price[0].source_locator == cell_locator("Demo经营资料.xlsx", "商品参数", 5, 2)
    assert price[0].status is EvidenceStatus.PENDING
    assert price[0].source_hash == "hash"

    review = [item for item in evidences if item.field_id == "excel.用户评价.J01.评价文本"]
    assert len(review) == 1
    assert "坐下不勒" in str(review[0].value)
    assert review[0].source_locator == cell_locator("Demo经营资料.xlsx", "用户评价", 5, 3)


def test_txt_and_md_extractors_emit_line_and_heading_locators(tmp_path: Path) -> None:
    brand_path = tmp_path / "品牌规则.txt"
    brand_path.write_text("品牌名称：界线 JIEXIAN\n禁用词：显瘦神裤\n", encoding="utf-8")
    md_path = tmp_path / "新品内容包.md"
    md_path.write_text(
        "## 商品理解卡\n- 建议价格：329 元\n\n### 用户真正喜欢\n- 坐下不勒\n",
        encoding="utf-8",
    )

    brand_source = _source(
        tmp_path,
        "品牌规则.txt",
        source_kind=SourceKind.BRAND_RULE,
        file_type="txt",
    )
    md_source = _source(
        tmp_path,
        "新品内容包.md",
        source_kind=SourceKind.UNKNOWN,
        file_type="md",
    )

    brand_items = extract_evidence_for_source(brand_source, brand_path)
    md_items = extract_evidence_for_source(md_source, md_path)

    assert brand_items[0].source_locator == line_locator("品牌规则.txt", 1)
    assert brand_items[0].value == "界线 JIEXIAN"
    assert any(item.value == "329 元" for item in md_items)
    assert any("商品理解卡" in item.source_locator for item in md_items)


def test_extract_from_index_skips_binary_sources(tmp_path: Path) -> None:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "品牌规则.txt").write_text("品牌名称：测试品牌\n", encoding="utf-8")
    (pack / "商品图片").mkdir()
    (pack / "商品图片" / "01.png").write_bytes(b"png")

    report = scan_library(pack)
    extraction = extract_evidence_from_index(report.index(), report.library_root)

    assert extraction.evidences
    assert any(item.field_id == "excel.商品参数.建议价格" for item in extraction.evidences)
    assert any(skip.reason == "unsupported_source_kind" for skip in extraction.skipped)


def test_evidence_json_writes_outside_source_root(tmp_path: Path) -> None:
    pack = tmp_path / "pack"
    output = tmp_path / "workspace"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")

    report = scan_library(pack)
    extraction = extract_evidence_from_index(report.index(), report.library_root)
    errors = extraction.write_json(output, pack)

    assert errors == []
    written = output / extraction.product_id / EVIDENCE_FILE_NAME
    assert written.is_file()
    assert "evidences" in written.read_text(encoding="utf-8")


def test_missing_xlsx_sheet_raises(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "商品参数"
    workbook.save(path)
    source = _source(tmp_path, "broken.xlsx")

    with pytest.raises(ValueError, match="missing_sheets"):
        extract_evidence_for_source(source, path)
