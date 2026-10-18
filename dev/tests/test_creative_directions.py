from __future__ import annotations

from pathlib import Path

import openpyxl

from app.core.business_judgment import build_business_judgment
from app.core.creative_directions import (
    DIRECTIONS_FILE_NAME,
    build_creative_directions,
    direction_id_of,
    evidence_kinds_for_direction,
)
from app.core.evidence_extractor import extract_evidence_from_index
from app.core.product_understanding import build_product_understanding
from app.core.source_scanner import scan_library
from app.steps.extract_evidence import ExtractEvidenceStep
from app.steps.generate_directions import GenerateDirectionsStep
from app.steps.judge_business import JudgeBusinessStep
from app.steps.pipeline import Pipeline
from app.steps.scan_sources import ScanSourcesStep
from app.steps.understand_product import UnderstandProductStep


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    product_sheet = workbook.active
    product_sheet.title = "商品参数"
    product_sheet.append(["商品参数"])
    product_sheet.append(["演示数据"])
    product_sheet.append([])
    product_sheet.append(["字段", "值", "分类", "数据使用提示"])
    product_sheet.append(["商品名称", "测试牛仔裤 01 中靛蓝", "商品", ""])
    product_sheet.append(["建议价格", 329, "价格", ""])
    product_sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    product_sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    product_sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    product_sheet.append(["主要场景", "通勤久坐 咖啡馆", "场景", ""])
    product_sheet.append(["版型", "高腰直筒", "版型", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐一上午腰头也不会勒", "正向", "腰部", "办公室"])
    reviews.append(["J03", 5, "中靛蓝上班穿有分寸", "正向", "中靛蓝", "通勤"])
    reviews.append(["J13", 3, "担心不同身高裤脚会堆", "负向", "裤长", "通勤"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "显瘦直筒裤", "显瘦 小个子", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤 闭眼入", "禁入文案"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "历史内容.txt").write_text(
        "结论： “显瘦”出现 4 次；“百搭”出现 3 次。\n",
        encoding="utf-8",
    )
    images = pack / "商品图片"
    images.mkdir()
    (images / "01_牛仔裤平铺主图.png").write_bytes(b"png")
    (images / "02_正面上身图.png").write_bytes(b"png2")
    (images / "03_腰部细节图.png").write_bytes(b"png3")
    (images / "04_久坐咖啡馆场景图.png").write_bytes(b"png4")
    return pack


def test_direction_id_is_stable() -> None:
    first = direction_id_of("prod_a", "方向 A")
    second = direction_id_of("prod_a", "方向 A")
    third = direction_id_of("prod_a", "方向 B")
    assert first == second
    assert first != third


def test_build_creative_directions_produces_three_distinct_directions(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())
    judgment = build_business_judgment(understanding, extraction)
    report = build_creative_directions(judgment, understanding, extraction)

    assert len(report.directions) == 3
    names = [item.name for item in report.directions]
    assert len(set(names)) == 3
    assert "腰里也有余量" in names[0]
    assert "咖啡馆" in names[1]
    assert "三档裤长" in names[2]

    for direction in report.directions:
        assert direction.risks
        assert direction.open_questions
        assert direction.asset_ids
        assert direction.avoid_expression
        kinds = evidence_kinds_for_direction(direction, extraction.evidences)
        assert len(kinds) >= 2


def test_generate_directions_step_runs_full_pipeline(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    context = {
        "material_root": pack,
        "output_root": output,
    }
    results = Pipeline(
        [
            ScanSourcesStep(),
            ExtractEvidenceStep(),
            UnderstandProductStep(),
            JudgeBusinessStep(),
            GenerateDirectionsStep(),
        ]
    ).run(context)

    assert [result.step_name for result in results][-1] == "generate_directions"
    report = results[-1].data["report"]
    assert len(report.directions) == 3

    written = output / results[0].data["product_id"] / DIRECTIONS_FILE_NAME
    assert written.is_file()
