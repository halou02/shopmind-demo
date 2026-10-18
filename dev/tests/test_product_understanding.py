from __future__ import annotations

from pathlib import Path

import openpyxl

from app.core.evidence_extractor import extract_evidence_from_index
from app.core.models import EvidenceStatus
from app.core.product_understanding import (
    SECTION_BRAND_MEMORY,
    SECTION_DATA_GAPS,
    SECTION_PRODUCT_FACTS,
    SECTION_REVIEW_INSIGHTS,
    UNDERSTANDING_FILE_NAME,
    apply_product_form,
    build_product_understanding,
    fact_id_of,
)
from app.core.source_scanner import scan_library
from app.steps.pipeline import Pipeline
from app.steps.scan_sources import ScanSourcesStep
from app.steps.extract_evidence import ExtractEvidenceStep
from app.steps.understand_product import UnderstandProductStep


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    product_sheet = workbook.active
    product_sheet.title = "商品参数"
    product_sheet.append(["商品参数"])
    product_sheet.append(["演示数据"])
    product_sheet.append([])
    product_sheet.append(["字段", "值", "分类", "数据使用提示"])
    product_sheet.append(["商品名称", "测试牛仔裤 01", "商品", ""])
    product_sheet.append(["建议价格", 329, "价格", ""])
    product_sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    product_sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    product_sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "腰头有支撑，坐下不勒", "正向", "腰部", "办公室"])
    reviews.append(["J02", 4, "尺码指引不够细", "负向", "尺码", "通勤"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "显瘦直筒裤", "显瘦 小个子", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤 闭眼入", "禁入文案"])
    brand.append(["历史高频词", "百搭 显瘦", "不可做主标题"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "品牌规则.txt").write_text(
        "品牌定位：测试品牌定位\n禁用词：显瘦神裤\n历史表达提醒：百搭已高频\n",
        encoding="utf-8",
    )
    (pack / "历史内容.txt").write_text(
        "2026-04-18  晴日直筒裤\n标题：显瘦直筒裤\n结论： “显瘦”出现 4 次。\n",
        encoding="utf-8",
    )
    (pack / "商品图片").mkdir()
    (pack / "商品图片" / "01_平铺主图.png").write_bytes(b"png")
    return pack


def test_fact_id_is_stable() -> None:
    first = fact_id_of("prod_a", "product_facts", "建议价格")
    second = fact_id_of("prod_a", "product_facts", "建议价格")
    third = fact_id_of("prod_b", "product_facts", "建议价格")

    assert first == second
    assert first != third


def test_build_product_understanding_groups_facts_and_gaps(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())

    assert understanding.product_name == "测试牛仔裤 01"
    assert understanding.boundary_notice == "仅基于本地资料"
    assert understanding.not_connected_notice == "未接入，不作判断"

    product_facts = understanding.facts_for_section(SECTION_PRODUCT_FACTS)
    assert any("329 元" in fact.summary for fact in product_facts)
    assert all(fact.evidence_ids for fact in product_facts)

    review_facts = understanding.facts_for_section(SECTION_REVIEW_INSIGHTS)
    assert any(fact.label == "用户真正喜欢" for fact in review_facts)
    assert any(fact.label == "用户主要顾虑" for fact in review_facts)
    assert any("J01" in fact.summary for fact in review_facts)

    brand_facts = understanding.facts_for_section(SECTION_BRAND_MEMORY)
    assert any("显瘦神裤" in fact.summary for fact in brand_facts)

    gap_section = understanding.section(SECTION_DATA_GAPS)
    assert gap_section is not None
    assert gap_section.gaps
    assert all(fact.is_gap for fact in gap_section.facts)

    assert len(understanding.assets) == 1
    assert understanding.assets[0].role == "main_image"


def test_understand_product_step_runs_after_extract_evidence(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    context = {
        "material_root": pack,
        "output_root": output,
    }
    results = Pipeline(
        [ScanSourcesStep(), ExtractEvidenceStep(), UnderstandProductStep()]
    ).run(context)

    assert [result.step_name for result in results] == [
        "scan_sources",
        "extract_evidence",
        "understand_product",
    ]
    assert results[-1].data["product_name"] == "测试牛仔裤 01"
    assert results[-1].data["gap_count"] >= 1

    written = output / results[0].data["product_id"] / UNDERSTANDING_FILE_NAME
    assert written.is_file()


def test_apply_product_form_overrides_understanding(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())
    merged = apply_product_form(
        understanding,
        {
            "product_name": "轻序托特包",
            "product_category": "箱包",
            "product_price": "120",
            "product_desc": "绿色，容量大",
            "product_brand": "TMM",
        },
    )

    assert merged.product_name == "轻序托特包"
    facts = {fact.label: fact.summary for fact in merged.facts_for_section(SECTION_PRODUCT_FACTS)}
    assert facts["商品类别"] == "箱包"
    assert facts["建议价格"] == "120 元"
    assert facts["主要场景"] == "绿色，容量大"
    assert facts["品牌名称"] == "TMM"


def test_understanding_facts_default_to_pending_status(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())

    for section in understanding.sections:
        for fact in section.facts:
            if not fact.is_gap:
                assert fact.status is EvidenceStatus.PENDING
