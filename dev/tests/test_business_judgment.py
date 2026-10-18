from __future__ import annotations

from pathlib import Path

import openpyxl

from app.core.business_judgment import (
    JUDGMENT_FILE_NAME,
    build_business_judgment,
    judgment_id_of,
)
from app.core.evidence_extractor import extract_evidence_from_index
from app.core.product_understanding import build_product_understanding
from app.core.source_scanner import scan_library
from app.steps.extract_evidence import ExtractEvidenceStep
from app.steps.judge_business import JudgeBusinessStep
from app.steps.pipeline import Pipeline
from app.steps.scan_sources import ScanSourcesStep
from app.steps.understand_product import UnderstandProductStep


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    product_sheet = workbook.active
    product_sheet.title = "商品参数"
    product_sheet.append(["商品参数"])
    product_sheet.append(["演示数据"])
    product_sheet.append([])
    product_sheet.append(["字段", "值", "分类", "数据使用提示"])
    product_sheet.append(["商品名称", "测试牛仔裤 01 中靛蓝", "商品", ""])
    product_sheet.append(["建议价格", 329, "价格", ""])
    product_sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    product_sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    product_sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    product_sheet.append(["主要场景", "通勤久坐 咖啡馆", "场景", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐一上午腰头也不会勒", "正向", "腰部", "办公室"])
    reviews.append(["J02", 5, "直筒腿不是贴着的，走路舒服", "正向", "直筒", "通勤"])
    reviews.append(["J11", 3, "尺码要看腰围和臀围", "负向", "尺码", "通勤"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "显瘦直筒裤", "显瘦 小个子", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["历史高频词", "百搭 显瘦", "不可做主标题"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "历史内容.txt").write_text(
        "结论： “显瘦”出现 4 次；“百搭”出现 3 次。\n",
        encoding="utf-8",
    )
    return pack


def test_judgment_id_is_stable() -> None:
    first = judgment_id_of("prod_a")
    second = judgment_id_of("prod_a")
    third = judgment_id_of("prod_b")
    assert first == second
    assert first != third


def test_build_business_judgment_uses_local_evidence(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())
    report = build_business_judgment(understanding, extraction)

    assert report.judgment.boundary_notice == "仅基于本地资料"
    assert "未接入，不作判断" in report.judgment.not_connected_notice
    assert "显瘦神裤" in report.judgment.opportunity
    assert "99% 棉 1% 氨纶" in report.product_evidence
    assert "J01" in report.user_evidence
    assert report.judgment.evidence_ids
    assert report.judgment.risks
    assert report.judgment.open_questions
    assert report.judgment.next_actions


def test_judge_business_step_runs_after_understand_product(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    context = {
        "material_root": pack,
        "output_root": output,
    }
    results = Pipeline(
        [
            ScanSourcesStep(),
            ExtractEvidenceStep(),
            UnderstandProductStep(),
            JudgeBusinessStep(),
        ]
    ).run(context)

    assert [result.step_name for result in results] == [
        "scan_sources",
        "extract_evidence",
        "understand_product",
        "judge_business",
    ]
    report = results[-1].data["report"]
    assert "久坐" in report.judgment.opportunity or "显瘦神裤" in report.judgment.opportunity

    written = output / results[0].data["product_id"] / JUDGMENT_FILE_NAME
    assert written.is_file()
