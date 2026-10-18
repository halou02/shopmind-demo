from __future__ import annotations

from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from app.web.server import app


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "商品参数"
    sheet.append(["商品参数"])
    sheet.append(["演示数据"])
    sheet.append([])
    sheet.append(["字段", "值", "分类", "数据使用提示"])
    sheet.append(["商品名称", "Web API 测试牛仔裤", "商品", ""])
    sheet.append(["建议价格", 329, "价格", ""])
    sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    sheet.append(["主要场景", "通勤久坐", "场景", ""])
    sheet.append(["版型", "高腰直筒", "版型", ""])
    sheet.append(["品牌禁用词", "显瘦神裤", "品牌", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐下不勒", "正向", "腰部", "办公室"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "标题", "显瘦", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤", "禁入文案"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "历史内容.txt").write_text("结论： “显瘦”出现 4 次。\n", encoding="utf-8")
    images = pack / "商品图片"
    images.mkdir()
    (images / "01_主图.png").write_bytes(b"png")
    return pack


def test_web_select_reference_images(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    pack = _build_pack(tmp_path)
    client = TestClient(app)
    preview = client.post("/api/material-preview", json={"material_path": str(pack)})
    assert preview.status_code == 200
    images = preview.json()["images"]
    assert images
    selected = client.post(
        "/api/select-reference-images",
        json={"relative_paths": [images[0]["relative_path"]]},
    )
    assert selected.status_code == 200
    payload = selected.json()
    assert payload["count"] == 1
    assert payload["reference_images"]


def test_web_upload_reference_files(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    client = TestClient(app)
    files = [
        ("files", ("brief.txt", b"green tote bag large capacity", "text/plain")),
    ]
    response = client.post("/api/upload-files?target=library", files=files)
    assert response.status_code == 200
    payload = response.json()
    assert payload["added_count"] == 1
    assert any("brief.txt" in item["name"] for item in payload["reference_files"])


def test_web_materials_confirm_after_upload(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    client = TestClient(app)
    files = [
        ("files", ("demo-1.png", b"\x89PNG\r\n\x1a\n", "image/png")),
    ]
    upload = client.post("/api/upload-files", files=files)
    assert upload.status_code == 200
    payload = upload.json()
    assert payload["materials_ready"] is True
    assert payload["reference_images"]
    confirm = client.post("/api/materials/confirm", json={})
    assert confirm.status_code == 200
    assert confirm.json()["materials_ready"] is True


def test_web_upload_images_and_analyze(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    client = TestClient(app)
    files = [
        ("files", ("demo-1.png", b"\x89PNG\r\n\x1a\n", "image/png")),
        ("files", ("demo-2.jpg", b"\xff\xd8\xff", "image/jpeg")),
    ]
    response = client.post("/api/upload-images", files=files)
    assert response.status_code == 200
    payload = response.json()
    assert payload["added_count"] == 2
    assert len(payload["images"]) >= 2
    assert len(payload["reference_images"]) >= 2
    assert Path(payload["material_path"]).is_dir()

    analysis = client.post(
        "/api/analyze",
        json={
            "material_path": payload["material_path"],
            "workspace_path": str(tmp_path / "workspace"),
            "product_form": {"product_name": "图片测试商品", "product_category": "箱包"},
        },
    )
    assert analysis.status_code == 200
    assert analysis.json()["ok"] is True


def test_web_health_and_static() -> None:
    client = TestClient(app)
    health = client.get("/api/health")
    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    index = client.get("/")
    assert index.status_code == 200
    assert "ShopMind" in index.text

    meta = client.get("/api/meta")
    assert meta.status_code == 200
    assert "help" in meta.json()


def test_web_analyze_and_publish(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    client = TestClient(app)

    preview = client.post("/api/material-preview", json={"material_path": str(pack)})
    assert preview.status_code == 200
    preview_payload = preview.json()
    assert len(preview_payload["images"]) >= 1
    assert preview_payload["stats"]["total_count"] >= 1

    form = client.post(
        "/api/product-form",
        json={"product_name": "测试商品", "product_category": "服装", "product_price": "329"},
    )
    assert form.status_code == 200

    analysis = client.post(
        "/api/analyze",
        json={
            "material_path": str(pack),
            "workspace_path": str(tmp_path / "workspace"),
            "product_form": {"product_name": "测试商品"},
        },
    )
    assert analysis.status_code == 200
    payload = analysis.json()
    assert payload["ok"] is True
    assert payload["directions"]["direction_count"] == 3

    notes = client.post("/api/understanding-notes", json={"notes": "人工修正：强调久坐舒适"})
    assert notes.status_code == 200

    regenerate = client.post("/api/regenerate-directions")
    assert regenerate.status_code == 200
    assert regenerate.json()["ok"] is True

    direction_id = payload["directions"]["directions"][0]["direction_id"]
    publish = client.post("/api/publish", json={"direction_id": direction_id})
    assert publish.status_code == 200
    publish_payload = publish.json()
    assert publish_payload["ok"] is True
    assert publish_payload["content"]["item_count"] == 6
    assert publish_payload["export"]["files"]

    files = client.get("/api/export/files")
    assert files.status_code == 200
    assert files.json()["files"]

    first_file = files.json()["files"][0]["relative_path"]
    download = client.get(f"/api/export/download?path={first_file}")
    assert download.status_code == 200

    zip_resp = client.get("/api/export/zip")
    assert zip_resp.status_code == 200
    assert zip_resp.headers["content-type"].startswith("application/zip")

    products = client.get("/api/products")
    assert products.status_code == 200
    assert products.json()["products"]

    projects = client.get("/api/projects")
    assert projects.status_code == 200
    assert projects.json()["projects"]

    state = client.get("/api/state")
    assert state.status_code == 200
    assert state.json()["analysis"] is not None
