from __future__ import annotations

from pathlib import Path

import openpyxl

from app.ui.pipeline_runner import run_analysis, run_publish


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "商品参数"
    sheet.append(["商品参数"])
    sheet.append(["演示数据"])
    sheet.append([])
    sheet.append(["字段", "值", "分类", "数据使用提示"])
    sheet.append(["商品名称", "UI 测试牛仔裤", "商品", ""])
    sheet.append(["建议价格", 329, "价格", ""])
    sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    sheet.append(["主要场景", "通勤久坐", "场景", ""])
    sheet.append(["版型", "高腰直筒", "版型", ""])
    sheet.append(["品牌禁用词", "显瘦神裤", "品牌", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐下不勒", "正向", "腰部", "办公室"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "标题", "显瘦", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤", "禁入文案"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "历史内容.txt").write_text("结论： “显瘦”出现 4 次。\n", encoding="utf-8")
    images = pack / "商品图片"
    images.mkdir()
    (images / "01_主图.png").write_bytes(b"png")
    (images / "02_上身.png").write_bytes(b"png2")
    (images / "03_细节.png").write_bytes(b"png3")
    (images / "04_场景.png").write_bytes(b"png4")
    return pack


def test_run_analysis_and_publish(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    workspace = tmp_path / "workspace"
    analysis = run_analysis(pack, workspace)
    assert analysis.ok
    assert analysis.understanding is not None
    assert analysis.judgment is not None
    assert analysis.directions is not None
    assert len(analysis.directions.directions) == 3

    direction_id = analysis.directions.directions[0].direction_id
    publish = run_publish(analysis.context, selected_direction_id=direction_id, workspace_root=workspace)
    assert publish.ok
    assert publish.content is not None
    assert publish.export is not None
    assert publish.export.package_root.is_dir()
    assert len(publish.content.items) == 6
