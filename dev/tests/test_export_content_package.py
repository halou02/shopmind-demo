from __future__ import annotations

from pathlib import Path

import openpyxl

from app.core.business_judgment import build_business_judgment
from app.core.content_generator import ContentGenerationReport, build_content_package
from app.core.creative_directions import build_creative_directions
from app.core.evidence_extractor import extract_evidence_from_index
from app.core.models import ContentItem, ValidationLevel
from app.core.product_understanding import build_product_understanding
from app.core.source_scanner import scan_library
from app.export.content_package import EXPORT_STATUS_BLOCKED, export_content_package
from app.export.manifest import verify_manifest
from app.steps.export_content_package import ExportContentPackageStep
from app.steps.extract_evidence import ExtractEvidenceStep
from app.steps.generate_content import GenerateContentStep
from app.steps.generate_directions import GenerateDirectionsStep
from app.steps.judge_business import JudgeBusinessStep
from app.steps.pipeline import Pipeline
from app.steps.scan_sources import ScanSourcesStep
from app.steps.understand_product import UnderstandProductStep


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    product_sheet = workbook.active
    product_sheet.title = "商品参数"
    product_sheet.append(["商品参数"])
    product_sheet.append(["演示数据"])
    product_sheet.append([])
    product_sheet.append(["字段", "值", "分类", "数据使用提示"])
    product_sheet.append(["商品名称", "测试牛仔裤 01 中靛蓝", "商品", ""])
    product_sheet.append(["品牌", "界线 JIEXIAN", "品牌", ""])
    product_sheet.append(["建议价格", 329, "价格", ""])
    product_sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    product_sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    product_sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    product_sheet.append(["主要场景", "通勤久坐 咖啡馆", "场景", ""])
    product_sheet.append(["版型", "高腰直筒", "版型", ""])
    product_sheet.append(["品牌禁用词", "显瘦神裤 闭眼入", "品牌", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐一上午腰头也不会勒", "正向", "腰部", "办公室"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "显瘦直筒裤", "显瘦 小个子", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤 闭眼入", "禁入文案"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "历史内容.txt").write_text(
        "结论： “显瘦”出现 4 次；“百搭”出现 3 次。\n",
        encoding="utf-8",
    )
    images = pack / "商品图片"
    images.mkdir()
    (images / "01_牛仔裤平铺主图.png").write_bytes(b"png")
    (images / "02_正面上身图.png").write_bytes(b"png2")
    (images / "03_腰部细节图.png").write_bytes(b"png3")
    (images / "04_久坐咖啡馆场景图.png").write_bytes(b"png4")
    return pack


def _full_reports(pack: Path, output: Path):
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())
    judgment = build_business_judgment(understanding, extraction)
    directions = build_creative_directions(judgment, understanding, extraction)
    content = build_content_package(directions, understanding, extraction)
    export = export_content_package(
        understanding=understanding,
        judgment_report=judgment,
        directions_report=directions,
        content_report=content,
        extraction=extraction,
        index=scan_report.index(),
        output_root=output,
        library_root=scan_report.library_root,
    )
    return scan_report, content, export


def test_export_content_package_writes_manifest_and_markdown(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    _, _, export = _full_reports(pack, output)

    assert export.export_status in {"ready", "partial"}
    assert export.package_root.is_dir()
    assert (export.package_root / "manifest.json").is_file()
    assert (export.package_root / "README.md").is_file()
    assert (export.package_root / "content" / "xiaohongshu.md").is_file()
    assert verify_manifest(export.package_root, export.manifest) == []


def test_export_is_blocked_when_content_validation_fails(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())
    judgment = build_business_judgment(understanding, extraction)
    directions = build_creative_directions(judgment, understanding, extraction)
    direction = directions.directions[0]
    bad_item = ContentItem(
        content_id="cnt_bad",
        product_id=direction.product_id,
        direction_id=direction.direction_id,
        channel="xiaohongshu",
        title="坏内容",
        body="显瘦神裤必买",
        evidence_ids=direction.evidence_ids,
        asset_ids=direction.asset_ids,
        validation_level=ValidationLevel.FAIL,
    )
    content = ContentGenerationReport(
        product_id=directions.product_id,
        product_name=understanding.product_name,
        direction_id=direction.direction_id,
        direction_name=direction.name,
        items=(bad_item,),
        validations=(),
    )
    export = export_content_package(
        understanding=understanding,
        judgment_report=judgment,
        directions_report=directions,
        content_report=content,
        extraction=extraction,
        index=scan_report.index(),
        output_root=output,
        library_root=scan_report.library_root,
    )
    assert export.export_status == EXPORT_STATUS_BLOCKED
    assert not export.package_root.exists()


def test_export_content_package_step_runs_full_pipeline(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    results = Pipeline(
        [
            ScanSourcesStep(),
            ExtractEvidenceStep(),
            UnderstandProductStep(),
            JudgeBusinessStep(),
            GenerateDirectionsStep(),
            GenerateContentStep(),
            ExportContentPackageStep(),
        ]
    ).run({"material_root": pack, "output_root": output})

    assert results[-1].step_name == "export_content_package"
    assert results[-1].data["export_status"] in {"ready", "partial"}
    package_root = Path(results[-1].data["package_root"])
    assert package_root.is_dir()
    assert (package_root / "manifest.json").is_file()
