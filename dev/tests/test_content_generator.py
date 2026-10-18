from __future__ import annotations

from pathlib import Path

import openpyxl

from app.core.business_judgment import build_business_judgment
from app.core.content_generator import (
    CHANNELS,
    CONTENT_FILE_NAME,
    build_content_package,
    content_id_of,
    validate_content_item,
)
from app.core.creative_directions import build_creative_directions
from app.core.evidence_extractor import extract_evidence_from_index
from app.core.models import ContentItem, ValidationLevel
from app.core.product_understanding import build_product_understanding
from app.core.source_scanner import scan_library
from app.steps.extract_evidence import ExtractEvidenceStep
from app.steps.generate_content import GenerateContentStep
from app.steps.generate_directions import GenerateDirectionsStep
from app.steps.judge_business import JudgeBusinessStep
from app.steps.pipeline import Pipeline
from app.steps.scan_sources import ScanSourcesStep
from app.steps.understand_product import UnderstandProductStep


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    product_sheet = workbook.active
    product_sheet.title = "商品参数"
    product_sheet.append(["商品参数"])
    product_sheet.append(["演示数据"])
    product_sheet.append([])
    product_sheet.append(["字段", "值", "分类", "数据使用提示"])
    product_sheet.append(["商品名称", "测试牛仔裤 01 中靛蓝", "商品", ""])
    product_sheet.append(["品牌", "界线 JIEXIAN", "品牌", ""])
    product_sheet.append(["建议价格", 329, "价格", ""])
    product_sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    product_sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    product_sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    product_sheet.append(["主要场景", "通勤久坐 咖啡馆", "场景", ""])
    product_sheet.append(["版型", "高腰直筒", "版型", ""])
    product_sheet.append(["品牌禁用词", "显瘦神裤 闭眼入", "品牌", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐一上午腰头也不会勒", "正向", "腰部", "办公室"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "显瘦直筒裤", "显瘦 小个子", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤 闭眼入", "禁入文案"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    (pack / "历史内容.txt").write_text(
        "结论： “显瘦”出现 4 次；“百搭”出现 3 次。\n",
        encoding="utf-8",
    )
    images = pack / "商品图片"
    images.mkdir()
    (images / "01_牛仔裤平铺主图.png").write_bytes(b"png")
    (images / "02_正面上身图.png").write_bytes(b"png2")
    (images / "03_腰部细节图.png").write_bytes(b"png3")
    (images / "04_久坐咖啡馆场景图.png").write_bytes(b"png4")
    return pack


def _reports(pack: Path):
    scan_report = scan_library(pack)
    extraction = extract_evidence_from_index(scan_report.index(), scan_report.library_root)
    understanding = build_product_understanding(extraction, scan_report.index())
    judgment = build_business_judgment(understanding, extraction)
    directions = build_creative_directions(judgment, understanding, extraction)
    return directions, understanding, extraction


def test_content_id_is_stable() -> None:
    first = content_id_of("prod_a", "dir_a", "xiaohongshu")
    second = content_id_of("prod_a", "dir_a", "xiaohongshu")
    third = content_id_of("prod_a", "dir_a", "douyin")
    assert first == second
    assert first != third


def test_build_content_package_generates_four_channels(tmp_path: Path) -> None:
    directions, understanding, extraction = _reports(_build_pack(tmp_path))
    report = build_content_package(directions, understanding, extraction)

    assert len(report.items) == len(CHANNELS)
    assert {item.channel for item in report.items} == set(CHANNELS)
    assert all(item.evidence_ids for item in report.items)
    assert all(item.asset_ids for item in report.items)
    assert all("仅基于本地资料" in item.body for item in report.items)
    assert all(item.validation_level is not ValidationLevel.FAIL for item in report.items)


def test_switching_direction_changes_generated_content(tmp_path: Path) -> None:
    directions, understanding, extraction = _reports(_build_pack(tmp_path))
    first = build_content_package(
        directions,
        understanding,
        extraction,
        selected_direction_id=directions.directions[0].direction_id,
    )
    second = build_content_package(
        directions,
        understanding,
        extraction,
        selected_direction_id=directions.directions[1].direction_id,
    )
    first_body = next(item.body for item in first.items if item.channel == "xiaohongshu")
    second_body = next(item.body for item in second.items if item.channel == "xiaohongshu")
    assert first_body != second_body


def test_validation_fails_on_forbidden_word_and_outside_evidence(tmp_path: Path) -> None:
    directions, understanding, extraction = _reports(_build_pack(tmp_path))
    direction = directions.directions[0]
    bad_item = ContentItem(
        content_id="cnt_test",
        product_id=direction.product_id,
        direction_id=direction.direction_id,
        channel="xiaohongshu",
        title="测试",
        body="这是显瘦神裤文案",
        evidence_ids=("outside_evidence",),
        asset_ids=direction.asset_ids,
    )
    level, results = validate_content_item(
        bad_item,
        direction=direction,
        banned_words=frozenset({"显瘦神裤"}),
        evidence_by_id={item.evidence_id: item for item in extraction.evidences},
    )
    assert level is ValidationLevel.FAIL
    assert any(result.code == "forbidden_word" for result in results)
    assert any(result.code == "evidence_outside_direction" for result in results)


def test_generate_content_step_runs_full_pipeline(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    output = tmp_path / "workspace"
    context = {
        "material_root": pack,
        "output_root": output,
    }
    results = Pipeline(
        [
            ScanSourcesStep(),
            ExtractEvidenceStep(),
            UnderstandProductStep(),
            JudgeBusinessStep(),
            GenerateDirectionsStep(),
            GenerateContentStep(),
        ]
    ).run(context)

    assert results[-1].step_name == "generate_content"
    assert results[-1].data["channel_count"] == len(CHANNELS)
    written = output / results[0].data["product_id"] / CONTENT_FILE_NAME
    assert written.is_file()
