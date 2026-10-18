from __future__ import annotations

import time
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from app.web.server import app


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "商品参数"
    sheet.append(["商品参数"])
    sheet.append(["演示数据"])
    sheet.append([])
    sheet.append(["字段", "值", "分类", "数据使用提示"])
    sheet.append(["商品名称", "媒体测试商品", "商品", ""])
    sheet.append(["建议价格", 329, "价格", ""])
    sheet.append(["材质", "99% 棉 1% 氨纶", "商品参数", ""])
    sheet.append(["裤长", "三档裤长 98 cm 102 cm 106 cm", "尺码", ""])
    sheet.append(["腰部活动余量", "约 2 cm", "商品参数", ""])
    sheet.append(["主要场景", "通勤久坐", "场景", ""])
    sheet.append(["版型", "高腰直筒", "版型", ""])
    sheet.append(["品牌禁用词", "显瘦神裤", "品牌", ""])

    reviews = workbook.create_sheet("用户评价")
    reviews.append(["用户评价"])
    reviews.append(["演示数据"])
    reviews.append([])
    reviews.append(["编号", "评分", "评价文本", "情绪", "标签", "使用场景"])
    reviews.append(["J01", 5, "坐下不勒", "正向", "腰部", "办公室"])

    history = workbook.create_sheet("历史内容")
    history.append(["历史内容"])
    history.append(["演示数据"])
    history.append([])
    history.append(["日期", "历史商品", "内容标题", "高频卖点", "重复表达风险"])
    history.append(["2026-04-18", "晴日直筒裤", "标题", "显瘦", "显瘦"])

    brand = workbook.create_sheet("品牌记忆")
    brand.append(["品牌记忆"])
    brand.append(["演示数据"])
    brand.append([])
    brand.append(["项目", "内容", "创作要求"])
    brand.append(["禁用词", "显瘦神裤", "禁入文案"])
    workbook.save(path)


def _build_pack(tmp_path: Path) -> Path:
    pack = tmp_path / "pack"
    pack.mkdir()
    _write_xlsx(pack / "Demo经营资料.xlsx")
    images = pack / "商品图片"
    images.mkdir()
    (images / "01_主图.png").write_bytes(b"png")
    (pack / "历史内容.txt").write_text("结论： “显瘦”出现 4 次。\n", encoding="utf-8")
    return pack


def test_generate_image_prompts_and_history(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    client = TestClient(app)

    preview = client.post("/api/material-preview", json={"material_path": str(pack)})
    images = preview.json()["images"]
    client.post("/api/select-reference-images", json={"relative_paths": [images[0]["relative_path"]]})

    analysis = client.post(
        "/api/analyze",
        json={
            "material_path": str(pack),
            "workspace_path": str(tmp_path / "workspace"),
            "product_form": {"product_name": "媒体测试商品"},
        },
    )
    direction_id = analysis.json()["directions"]["directions"][0]["direction_id"]
    publish = client.post("/api/publish", json={"direction_id": direction_id})
    assert publish.status_code == 200
    assert publish.json()["ok"] is True

    prompts = client.post("/api/generate-image-prompts")
    assert prompts.status_code == 200
    payload = prompts.json()
    assert payload["ok"] is True
    assert len(payload["image_prompts"]) >= 1

    media = client.post("/api/generate-media", json={"generate_image": True, "generate_video": False})
    assert media.status_code == 200
    payload = media.json()
    assert payload["image_prompts"]
    if payload.get("status") == "running":
        for _ in range(100):
            job = client.get("/api/media-job").json()
            if job.get("status") == "done":
                payload = job
                break
            if job.get("status") == "error":
                raise AssertionError(job.get("message"))
            time.sleep(0.05)
    assert payload.get("generated_media", {}).get("images")

    history = client.get("/api/history")
    assert history.status_code == 200
    items = history.json()["items"]
    assert items
    assert items[0]["product_name"] == "媒体测试商品"
    assert items[0]["has_image_prompts"] is True

    detail = client.get(f"/api/history/{items[0]['id']}")
    assert detail.status_code == 200
    assert detail.json()["direction_id"] == direction_id


def test_reference_images_add_remove(tmp_path: Path) -> None:
    pack = _build_pack(tmp_path)
    client = TestClient(app)
    preview = client.post("/api/material-preview", json={"material_path": str(pack)})
    images = preview.json()["images"]
    rel = images[0]["relative_path"]

    selected = client.post("/api/select-reference-images", json={"relative_paths": [rel]})
    assert selected.json()["count"] == 1

    removed = client.post("/api/reference-images/remove", json={"relative_path": rel})
    assert removed.status_code == 200
    assert removed.json()["count"] == 0
    assert not removed.json()["images"]

    added = client.post("/api/reference-images/add", json={"relative_paths": [rel]})
    assert added.status_code == 200
    assert added.json()["count"] == 0
